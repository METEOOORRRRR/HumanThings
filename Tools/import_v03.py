"""Extract the user's approved content tables without paraphrasing archive text."""
import json
import sys
from pathlib import Path
from docx import Document

doc = Document(sys.argv[1])
headers = [p.text.split(' — ', 1) for p in doc.paragraphs if ' — ' in p.text and p.text.startswith(('COM_', 'RES_', 'IND_', 'CUL_'))]
english = [
    'Violin', 'Drum Set', 'Piano', 'Frying Pan', 'Chopsticks', 'Cup Noodles', 'Smartphone', 'Camera', 'Radio',
    'Doll', 'Board Game', 'Toy Car', 'Hoodie', 'Sneakers', 'Umbrella', 'Comb', 'Mirror', 'Alarm Clock',
    'Hammer', 'Electric Drill', 'Tape Measure', 'Steering Wheel', 'Car Jack', 'Tire', 'Flyer', 'Sticky Note', 'Receipt',
    'Thermometer', 'Syringe', 'Stethoscope', 'Microscope', 'Test Tube', 'Safety Goggles', 'Walkie Talkie', 'Compass', 'Weather Station',
    'Novel', 'Map', 'Diary', 'Pencil', 'Report Card', 'Child Drawing', 'Portrait', 'Performance Ticket', 'Family Photograph',
]
locations, artifacts = [], []
for i, (identifier, name) in enumerate(headers):
    info = {r.cells[0].text: r.cells[1].text for r in doc.tables[18 + i * 2].rows[1:]}
    location = dict(id=identifier, zone=i//3, displayNameKnown=name, displayNameUnknown=info['초기 장소 설명'], moduleName=info['복원 모듈'], artifactIds=[])
    for row in doc.tables[19 + i * 2].rows[1:]:
        v = [c.text for c in row.cells]
        clues = [s.strip() for s in v[5].split(';')]
        artifact = dict(id=v[0], locationId=identifier, trueNameKo=v[1], trueNameEn=english[len(artifacts)], categoryTag=v[2], materialHint=v[3], functionHint=v[4], facilityHint=info['초기 장소 설명'], level1Clues=[clues[0]], level2Clues=[clues[min(1,len(clues)-1)], clues[min(2,len(clues)-1)]], archiveDescription=v[6], marsComment=v[7])
        artifacts.append(artifact)
        location['artifactIds'].append(v[0])
    locations.append(location)
assert len(locations) == 15 and len(artifacts) == 45
output = Path(sys.argv[2]); output.parent.mkdir(parents=True, exist_ok=True)
output.write_text(json.dumps(dict(locations=locations, artifacts=artifacts), ensure_ascii=False, indent=2), encoding='utf-8')
print('Imported 15 locations and 45 exact artifact records:', output)
